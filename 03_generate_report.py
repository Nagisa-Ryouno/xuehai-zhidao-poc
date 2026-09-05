# -*- coding: utf-8 -*-
"""
03_generate_report.py

学海智导 POC - 第三阶段

功能：

student_profiles.json
        +
learning_paths.json
        +
knowledge_points
        ↓
学生综合学习报告
        ↓
student_reports.json

本阶段主要完成：

1. 汇总学生基础信息
2. 汇总学生学情数据
3. 分析学生当前学习状态
4. 整理个性化学习路径
5. 生成学习阶段建议
6. 生成学习策略
7. 生成可直接供前端使用的 JSON 数据
"""

import json
from pathlib import Path

import openpyxl


# ============================================================
# 配置
# ============================================================

BASE_DIR = Path(__file__).resolve().parent

PROFILE_FILE = (
    BASE_DIR
    / "output"
    / "student_profiles.json"
)

LEARNING_PATH_FILE = (
    BASE_DIR
    / "output"
    / "learning_paths.json"
)

EXCEL_FILE = (
    BASE_DIR
    / "economics_learning_demo.xlsx"
)

OUTPUT_DIR = (
    BASE_DIR
    / "output"
)

OUTPUT_FILE = (
    OUTPUT_DIR
    / "student_reports.json"
)


# ============================================================
# 工具函数
# ============================================================

def safe_float(value, default=0):
    """安全转换为 float"""

    try:
        return float(value)

    except (
        TypeError,
        ValueError
    ):
        return default


def safe_int(value, default=0):
    """安全转换为 int"""

    try:
        return int(value)

    except (
        TypeError,
        ValueError
    ):
        return default


# ============================================================
# 页面标题
# ============================================================

print("=" * 70)
print("学海智导 · 学生综合学习报告生成器")
print("=" * 70)

print()
print("项目目录：", BASE_DIR)
print("学生画像：", PROFILE_FILE)
print("学习路径：", LEARNING_PATH_FILE)
print("Excel文件：", EXCEL_FILE)
print("报告输出：", OUTPUT_FILE)


# ============================================================
# 检查文件
# ============================================================

print()
print("正在检查输入文件...")

required_files = [
    PROFILE_FILE,
    LEARNING_PATH_FILE,
    EXCEL_FILE
]

for file_path in required_files:

    if not file_path.exists():

        print()
        print(
            f"❌ 文件不存在：{file_path}"
        )

        print()
        print(
            "请确认前面的步骤已经成功执行。"
        )

        raise SystemExit(1)


print("✓ 所有输入文件检查通过。")


# ============================================================
# 读取学生画像
# ============================================================

print()
print("正在读取学生画像...")

with open(
    PROFILE_FILE,
    "r",
    encoding="utf-8"
) as f:

    profiles = json.load(f)


print(
    f"✓ 成功读取 {len(profiles)} 名学生的学情画像。"
)


# ============================================================
# 读取学习路径
# ============================================================

print()
print("正在读取个性化学习路径...")

with open(
    LEARNING_PATH_FILE,
    "r",
    encoding="utf-8"
) as f:

    learning_paths = json.load(f)


print(
    f"✓ 成功读取 {len(learning_paths)} 名学生的学习路径。"
)


# ============================================================
# 读取知识点
# ============================================================

print()
print("正在读取知识点信息...")

wb = openpyxl.load_workbook(
    EXCEL_FILE,
    data_only=True
)

if "knowledge_points" not in wb.sheetnames:

    print(
        "❌ Excel 中不存在 knowledge_points 工作表"
    )

    raise SystemExit(1)


ws = wb["knowledge_points"]

knowledge_points = {}

for row in ws.iter_rows(
    min_row=2,
    values_only=True
):

    if not row:
        continue

    (
        knowledge_id,
        knowledge_name,
        chapter,
        description,
        difficulty,
        prerequisite,
        next_knowledge
    ) = row

    if knowledge_id is None:
        continue

    knowledge_id = str(
        knowledge_id
    ).strip()

    knowledge_points[
        knowledge_id
    ] = {

        "knowledge_id":
            knowledge_id,

        "knowledge_name":
            knowledge_name,

        "chapter":
            chapter,

        "description":
            description,

        "difficulty":
            safe_float(difficulty)
    }


print(
    f"✓ 成功读取 {len(knowledge_points)} 个知识点。"
)


# ============================================================
# 获取知识点
# ============================================================

def get_knowledge(
    knowledge_id
):

    return knowledge_points.get(
        knowledge_id,
        {
            "knowledge_id":
                knowledge_id,

            "knowledge_name":
                knowledge_id,

            "chapter":
                "",

            "description":
                "",

            "difficulty":
                0
        }
    )


# ============================================================
# 生成掌握程度说明
# ============================================================

def generate_mastery_description(
    average_accuracy
):

    accuracy = safe_float(
        average_accuracy
    )

    if accuracy >= 85:

        return (
            "整体知识掌握较好，"
            "已经具备较稳定的基础能力，"
            "后续应减少重复性基础训练，"
            "重点提升综合应用和知识迁移能力。"
        )

    elif accuracy >= 70:

        return (
            "整体知识掌握处于中等水平，"
            "已经建立基本知识框架，"
            "但部分知识点仍存在理解不牢的问题，"
            "适合进行针对性强化训练。"
        )

    elif accuracy >= 60:

        return (
            "整体知识掌握程度偏弱，"
            "部分核心知识点存在明显薄弱情况，"
            "建议结合知识依赖关系进行针对性补强。"
        )

    else:

        return (
            "整体知识掌握程度较弱，"
            "当前存在较多基础知识漏洞，"
            "建议优先重建基础知识体系，"
            "再逐步进入综合应用训练。"
        )


# ============================================================
# 生成学习行为说明
# ============================================================

def generate_behavior_description(
    profile_summary
):

    activity_level = profile_summary.get(
        "activity_level",
        ""
    )

    speed_status = profile_summary.get(
        "speed_status",
        ""
    )

    completion_level = profile_summary.get(
        "completion_level",
        ""
    )

    descriptions = []

    # 活跃度

    if activity_level == "高":

        descriptions.append(
            "学习参与度较高"
        )

    elif activity_level == "中":

        descriptions.append(
            "学习参与度处于中等水平"
        )

    else:

        descriptions.append(
            "学习参与度偏低"
        )

    # 答题速度

    if speed_status == "答题耗时较长":

        descriptions.append(
            "答题速度偏慢"
        )

    elif speed_status == "答题速度一般":

        descriptions.append(
            "答题速度一般"
        )

    else:

        descriptions.append(
            "答题速度较快"
        )

    # 完成度

    if completion_level == "高":

        descriptions.append(
            "学习任务完成度较高"
        )

    elif completion_level == "中":

        descriptions.append(
            "学习任务完成度处于中等水平"
        )

    else:

        descriptions.append(
            "学习任务完成度偏低"
        )

    return "，".join(
        descriptions
    ) + "。"


# ============================================================
# 生成核心问题
# ============================================================

def generate_main_problems(
    profile,
    learning_path_result
):

    problems = []

    overall = profile.get(
        "overall_profile",
        {}
    )

    weak_points = profile.get(
        "weak_knowledge_points",
        []
    )

    prerequisite_points = profile.get(
        "prerequisite_knowledge_points",
        []
    )

    average_accuracy = safe_float(
        overall.get(
            "average_accuracy"
        )
    )

    activity_level = overall.get(
        "activity_level",
        ""
    )

    speed_status = overall.get(
        "speed_status",
        ""
    )

    completion_level = overall.get(
        "completion_level",
        ""
    )

    # --------------------------------------------------------
    # 知识掌握
    # --------------------------------------------------------

    if average_accuracy < 60:

        problems.append(
            "整体知识基础较弱"
        )

    elif average_accuracy < 70:

        problems.append(
            "部分核心知识点掌握不足"
        )

    elif weak_points:

        problems.append(
            "存在局部知识薄弱点"
        )

    # --------------------------------------------------------
    # 前置知识
    # --------------------------------------------------------

    if prerequisite_points:

        problems.append(
            "部分知识薄弱与前置知识掌握不足有关"
        )

    # --------------------------------------------------------
    # 答题速度
    # --------------------------------------------------------

    if speed_status == "答题耗时较长":

        problems.append(
            "答题效率有待提升"
        )

    # --------------------------------------------------------
    # 学习完成度
    # --------------------------------------------------------

    if completion_level == "较低":

        problems.append(
            "学习任务完成度偏低"
        )

    # --------------------------------------------------------
    # 学习行为
    # --------------------------------------------------------

    if (
        activity_level == "高"
        and average_accuracy < 70
    ):

        problems.append(
            "学习参与度较高但学习效果仍需提升"
        )

    # --------------------------------------------------------
    # 如果没有明显问题
    # --------------------------------------------------------

    if not problems:

        problems.append(
            "当前没有明显的学习短板"
        )

    return problems


# ============================================================
# 生成学习阶段
# ============================================================

def generate_learning_stages(
    learning_path
):

    stages = []

    for item in learning_path:

        knowledge_id = item.get(
            "knowledge_id"
        )

        knowledge = get_knowledge(
            knowledge_id
        )

        stages.append({

            "stage":
                item.get(
                    "stage"
                ),

            "knowledge_id":
                knowledge_id,

            "knowledge_name":
                item.get(
                    "knowledge_name",
                    knowledge[
                        "knowledge_name"
                    ]
                ),

            "chapter":
                item.get(
                    "chapter",
                    knowledge[
                        "chapter"
                    ]
                ),

            "current_accuracy":
                safe_float(
                    item.get(
                        "current_accuracy"
                    )
                ),

            "priority":
                item.get(
                    "priority"
                ),

            "priority_score":
                safe_float(
                    item.get(
                        "priority_score"
                    )
                ),

            "learning_goal":
                item.get(
                    "learning_goal"
                ),

            "reason":
                item.get(
                    "reason"
                ),

            "difficulty":
                knowledge[
                    "difficulty"
                ],

            "description":
                knowledge[
                    "description"
                ]
        })

    return stages


# ============================================================
# 生成每日学习建议
# ============================================================

def generate_daily_plan(
    profile,
    learning_path_result
):

    overall = profile.get(
        "overall_profile",
        {}
    )

    average_accuracy = safe_float(
        overall.get(
            "average_accuracy"
        )
    )

    completion_level = overall.get(
        "completion_level",
        ""
    )

    recommendation_type = (
        learning_path_result.get(
            "recommendation_type",
            ""
        )
    )

    path_length = len(
        learning_path_result.get(
            "learning_path",
            []
        )
    )

    # --------------------------------------------------------
    # 默认
    # --------------------------------------------------------

    daily_minutes = 30

    question_count = 8

    focus = (
        "知识理解 + 基础练习"
    )

    # --------------------------------------------------------
    # 基础补强
    # --------------------------------------------------------

    if recommendation_type == "基础补强":

        daily_minutes = 45

        question_count = 10

        focus = (
            "基础知识理解 + 前置知识补齐"
        )

    # --------------------------------------------------------
    # 高参与度转化
    # --------------------------------------------------------

    elif recommendation_type == "高参与度转化训练":

        daily_minutes = 40

        question_count = 12

        focus = (
            "错题复盘 + 理解训练 + 练习转化"
        )

    # --------------------------------------------------------
    # 限时训练
    # --------------------------------------------------------

    elif (
        "限时训练"
        in recommendation_type
    ):

        daily_minutes = 40

        question_count = 15

        focus = (
            "综合应用 + 限时训练"
        )

    # --------------------------------------------------------
    # 能力提升
    # --------------------------------------------------------

    elif recommendation_type == "综合能力提升":

        daily_minutes = 35

        question_count = 10

        focus = (
            "综合应用 + 跨知识点训练"
        )

    # --------------------------------------------------------
    # 学习完成度低
    # --------------------------------------------------------

    if completion_level == "较低":

        daily_minutes = min(
            daily_minutes,
            30
        )

        question_count = min(
            question_count,
            8
        )

    # --------------------------------------------------------
    # 路径过长
    # --------------------------------------------------------

    if path_length >= 7:

        focus += (
            "，建议分阶段完成"
        )

    return {

        "recommended_minutes":
            daily_minutes,

        "recommended_questions":
            question_count,

        "focus":
            focus
    }


# ============================================================
# 生成 AI 风格学习总结
# ============================================================

def generate_ai_summary(
    profile,
    learning_path_result
):

    student = profile.get(
        "student",
        {}
    )

    overall = profile.get(
        "overall_profile",
        {}
    )

    average_accuracy = safe_float(
        overall.get(
            "average_accuracy"
        )
    )

    mastery_level = overall.get(
        "mastery_level",
        ""
    )

    recommendation_type = (
        learning_path_result.get(
            "recommendation_type",
            ""
        )
    )

    learning_path = (
        learning_path_result.get(
            "learning_path",
            []
        )
    )

    student_name = student.get(
        "student_name",
        "该学生"
    )

    if recommendation_type == "基础补强":

        summary = (
            f"{student_name}当前整体知识掌握程度为"
            f"{mastery_level}，平均正确率为"
            f"{average_accuracy:.2f}%。"
            f"当前主要问题是基础知识掌握不够扎实，"
            f"因此系统将学习策略定位为“基础补强”。"
        )

        if learning_path:

            first_point = learning_path[0].get(
                "knowledge_name",
                ""
            )

            summary += (
                f"建议首先从“{first_point}”开始，"
                f"逐步补齐基础知识链，"
                f"再进入后续综合应用训练。"
            )

        return summary

    if recommendation_type == "高参与度转化训练":

        return (
            f"{student_name}学习参与度较高，"
            f"但当前平均正确率为"
            f"{average_accuracy:.2f}%，"
            f"说明学习行为尚未完全转化为知识掌握。"
            f"建议通过错题复盘、概念理解和针对性练习，"
            f"提升学习行为向学习效果的转化效率。"
        )

    if recommendation_type == "综合能力提升":

        return (
            f"{student_name}当前平均正确率为"
            f"{average_accuracy:.2f}%，"
            f"整体知识掌握情况较好。"
            f"系统未发现需要优先补强的明显薄弱知识点，"
            f"因此建议减少基础重复训练，"
            f"重点开展综合应用、跨知识点训练和限时模拟。"
        )

    if (
        "限时训练"
        in recommendation_type
    ):

        return (
            f"{student_name}当前知识掌握情况总体较好，"
            f"但答题效率仍有提升空间。"
            f"建议在保持正确率的基础上，"
            f"增加限时训练和综合应用题，"
            f"逐步提升解题速度。"
        )

    return (
        f"{student_name}当前平均正确率为"
        f"{average_accuracy:.2f}%，"
        f"整体掌握程度为{mastery_level}。"
        f"系统识别出部分需要强化的知识点，"
        f"建议按照当前个性化学习路径进行针对性学习，"
        f"并根据后续学习效果动态调整学习计划。"
    )


# ============================================================
# 生成单个学生报告
# ============================================================

def generate_student_report(
    student_id,
    profile,
    learning_path_result
):

    student = profile.get(
        "student",
        {}
    )

    overall = profile.get(
        "overall_profile",
        {}
    )

    learning_path = (
        learning_path_result.get(
            "learning_path",
            []
        )
    )

    # --------------------------------------------------------
    # 学习阶段
    # --------------------------------------------------------

    stages = generate_learning_stages(
        learning_path
    )

    # --------------------------------------------------------
    # 核心问题
    # --------------------------------------------------------

    main_problems = generate_main_problems(
        profile,
        learning_path_result
    )

    # --------------------------------------------------------
    # 每日学习计划
    # --------------------------------------------------------

    daily_plan = generate_daily_plan(
        profile,
        learning_path_result
    )

    # --------------------------------------------------------
    # 学习行为
    # --------------------------------------------------------

    behavior_description = (
        generate_behavior_description(
            {
                "activity_level":
                    overall.get(
                        "activity_level"
                    ),

                "speed_status":
                    overall.get(
                        "speed_status"
                    ),

                "completion_level":
                    overall.get(
                        "completion_level"
                    )
            }
        )
    )

    # --------------------------------------------------------
    # AI总结
    # --------------------------------------------------------

    ai_summary = generate_ai_summary(
        profile,
        learning_path_result
    )

    # --------------------------------------------------------
    # 报告
    # --------------------------------------------------------

    return {

        "student_id":
            student_id,

        "student": {

            "student_id":
                student.get(
                    "student_id",
                    student_id
                ),

            "student_name":
                student.get(
                    "student_name",
                    ""
                ),

            "major":
                student.get(
                    "major",
                    ""
                ),

            "grade":
                student.get(
                    "grade",
                    ""
                ),

            "learning_goal":
                student.get(
                    "learning_goal",
                    ""
                )
        },

        # ====================================================
        # 学情概览
        # ====================================================

        "learning_overview": {

            "average_accuracy":
                safe_float(
                    overall.get(
                        "average_accuracy"
                    )
                ),

            "average_assessment_score":
                safe_float(
                    overall.get(
                        "average_assessment_score"
                    )
                ),

            "total_learning_time_minutes":
                safe_float(
                    overall.get(
                        "total_learning_time_minutes"
                    )
                ),

            "total_practice_count":
                safe_int(
                    overall.get(
                        "total_practice_count"
                    )
                ),

            "total_interaction_count":
                safe_int(
                    overall.get(
                        "total_interaction_count"
                    )
                ),

            "average_completion_rate":
                safe_float(
                    overall.get(
                        "average_completion_rate"
                    )
                ),

            "average_answer_time_seconds":
                safe_float(
                    overall.get(
                        "average_answer_time_seconds"
                    )
                ),

            "mastery_level":
                overall.get(
                    "mastery_level"
                ),

            "speed_status":
                overall.get(
                    "speed_status"
                ),

            "activity_level":
                overall.get(
                    "activity_level"
                ),

            "completion_level":
                overall.get(
                    "completion_level"
                )
        },

        # ====================================================
        # 学情诊断
        # ====================================================

        "diagnosis": {

            "mastery_description":
                generate_mastery_description(
                    overall.get(
                        "average_accuracy"
                    )
                ),

            "behavior_description":
                behavior_description,

            "main_problems":
                main_problems,

            "weak_knowledge_count":
                len(
                    profile.get(
                        "weak_knowledge_points",
                        []
                    )
                ),

            "prerequisite_knowledge_count":
                len(
                    profile.get(
                        "prerequisite_knowledge_points",
                        []
                    )
                )
        },

        # ====================================================
        # 个性化策略
        # ====================================================

        "personalized_strategy": {

            "recommendation_type":
                learning_path_result.get(
                    "recommendation_type"
                ),

            "strategy_description":
                ai_summary
        },

        # ====================================================
        # 学习路径
        # ====================================================

        "learning_path": stages,

        # ====================================================
        # 每日学习计划
        # ====================================================

        "daily_learning_plan":
            daily_plan,

        # ====================================================
        # 优化建议
        # ====================================================

        "optimization_suggestions":
            learning_path_result.get(
                "optimization_suggestion",
                []
            ),

        # ====================================================
        # AI学习总结
        # ====================================================

        "ai_summary":
            ai_summary
    }


# ============================================================
# 生成全部学生报告
# ============================================================

print()
print("=" * 70)
print("开始生成学生综合学习报告")
print("=" * 70)

student_reports = {}

for student_id, profile in profiles.items():

    print()
    print(
        f"正在生成 {student_id}..."
    )

    learning_path_result = (
        learning_paths.get(
            student_id,
            {}
        )
    )

    report = generate_student_report(
        student_id,
        profile,
        learning_path_result
    )

    student_reports[
        student_id
    ] = report

    student_name = report[
        "student"
    ].get(
        "student_name",
        ""
    )

    average_accuracy = report[
        "learning_overview"
    ].get(
        "average_accuracy",
        0
    )

    recommendation_type = report[
        "personalized_strategy"
    ].get(
        "recommendation_type",
        ""
    )

    path_length = len(
        report[
            "learning_path"
        ]
    )

    print(
        f"  学生：{student_name}"
    )

    print(
        f"  平均正确率："
        f"{average_accuracy:.2f}%"
    )

    print(
        f"  推荐策略："
        f"{recommendation_type}"
    )

    print(
        f"  学习路径长度："
        f"{path_length}"
    )


# ============================================================
# 保存 JSON
# ============================================================

print()
print("正在保存学生综合学习报告...")

OUTPUT_DIR.mkdir(
    exist_ok=True
)

with open(
    OUTPUT_FILE,
    "w",
    encoding="utf-8"
) as f:

    json.dump(
        student_reports,
        f,
        ensure_ascii=False,
        indent=2
    )


# ============================================================
# 最终输出
# ============================================================

print()
print("=" * 70)
print("学生综合学习报告生成完成！")
print("=" * 70)

print(
    f"学生数量："
    f"{len(student_reports)}"
)

print(
    f"输出文件："
    f"{OUTPUT_FILE}"
)

print()

for student_id, report in student_reports.items():

    student_name = report[
        "student"
    ].get(
        "student_name",
        ""
    )

    overview = report[
        "learning_overview"
    ]

    strategy = report[
        "personalized_strategy"
    ]

    diagnosis = report[
        "diagnosis"
    ]

    path = report[
        "learning_path"
    ]

    print(
        f"【{student_id} "
        f"{student_name}】"
    )

    print(
        f"掌握程度："
        f"{overview['mastery_level']}"
    )

    print(
        f"平均正确率："
        f"{overview['average_accuracy']:.2f}%"
    )

    print(
        f"推荐策略："
        f"{strategy['recommendation_type']}"
    )

    print(
        f"主要问题："
        f"{'；'.join(diagnosis['main_problems'])}"
    )

    print(
        f"学习路径："
        f"{len(path)} 个知识点"
    )

    print(
        f"每日建议："
        f"{report['daily_learning_plan']['recommended_minutes']}分钟"
        f" / "
        f"{report['daily_learning_plan']['recommended_questions']}道题"
    )

    print()
    print(
        "AI总结："
    )

    print(
        f"  {report['ai_summary']}"
    )

    print()
    print("-" * 70)

print()
print("✓ 第三阶段执行成功。")
print(
    "下一阶段可以将 student_reports.json "
    "接入 FastAPI 后端。"
)