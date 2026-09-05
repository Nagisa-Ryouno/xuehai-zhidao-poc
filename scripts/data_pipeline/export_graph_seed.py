# -*- coding: utf-8 -*-
"""
scripts/data_pipeline/export_graph_seed.py
Extracts knowledge points, prerequisite edges, and student historical records
from data/raw/economics_learning_demo.xlsx and generates data/seeds/knowledge_graph.json.
"""
import sys
from pathlib import Path
import json
import openpyxl

PROJECT_ROOT = Path(__file__).resolve().parent.parent.parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from app.core.config import settings

EXCEL_PATH = settings.EXCEL_RAW_FILE if hasattr(settings, "EXCEL_RAW_FILE") else PROJECT_ROOT / "data" / "raw" / "economics_learning_demo.xlsx"
OUTPUT_PATH = settings.KNOWLEDGE_GRAPH_FILE if hasattr(settings, "KNOWLEDGE_GRAPH_FILE") else PROJECT_ROOT / "data" / "seeds" / "knowledge_graph.json"

def split_ids(val):
    if val is None:
        return []
    s = str(val).strip()
    if not s or s.upper() in ("NULL", "NONE"):
        return []
    return [x.strip() for x in s.split(",") if x.strip()]

def export_knowledge_graph(excel_path: Path = None, output_path: Path = None):
    excel_file = excel_path or EXCEL_PATH
    out_file = output_path or OUTPUT_PATH

    if not excel_file.exists():
        raise FileNotFoundError(f"Excel file not found at {excel_file}")

    wb = openpyxl.load_workbook(excel_file, data_only=True)
    kps = {}
    edges = []
    records = {}

    # 1. Parse knowledge_points
    ws_kp = wb["knowledge_points"]
    for row in ws_kp.iter_rows(min_row=2, values_only=True):
        if not row[0]:
            continue
        k_id = str(row[0]).strip()
        k_name = str(row[1]).strip() if row[1] else ""
        chapter = str(row[2]).strip() if row[2] else ""
        description = str(row[3]).strip() if row[3] else ""
        difficulty = int(row[4]) if row[4] is not None else 1
        prereqs = split_ids(row[5])
        nexts = split_ids(row[6])

        kps[k_id] = {
            "knowledge_id": k_id,
            "knowledge_name": k_name,
            "chapter": chapter,
            "description": description,
            "difficulty": difficulty,
            "prerequisite": prereqs,
            "next_knowledge": nexts,
        }

        for p_id in prereqs:
            edges.append({
                "id": f"e-{p_id}-{k_id}",
                "source": p_id,
                "target": k_id,
                "type": "prerequisite",
            })

    # 2. Parse learning_records
    ws_rec = wb["learning_records"]
    for row in ws_rec.iter_rows(min_row=2, values_only=True):
        if not row[0] or not row[1] or not row[2]:
            continue
        sid = str(row[1]).strip()
        kid = str(row[2]).strip()
        acc = float(row[3]) if row[3] is not None else None
        avg_sec = float(row[4]) if row[4] is not None else None
        practice_cnt = int(row[5]) if row[5] is not None else None
        mistake_cnt = int(row[6]) if row[6] is not None else None
        duration_min = float(row[7]) if row[7] is not None else None
        score = float(row[10]) if len(row) > 10 and row[10] is not None else None

        if sid not in records:
            records[sid] = {}

        records[sid][kid] = {
            "accuracy": acc,
            "average_time_seconds": avg_sec,
            "practice_count": practice_cnt,
            "mistake_count": mistake_cnt,
            "duration_minutes": duration_min,
            "assessment_score": score,
        }

    output_data = {
        "metadata": {
            "version": "1.0",
            "source": "data/raw/economics_learning_demo.xlsx",
            "node_count": len(kps),
            "edge_count": len(edges),
        },
        "knowledge_points": kps,
        "edges": edges,
        "student_records": records,
    }

    out_file.parent.mkdir(parents=True, exist_ok=True)
    with open(out_file, "w", encoding="utf-8") as f:
        json.dump(output_data, f, ensure_ascii=False, indent=2)

    total_records = sum(len(v) for v in records.values())
    print(f"Exported {len(kps)} KPs, {len(edges)} edges, {total_records} records to {out_file}")
    return output_data

if __name__ == "__main__":
    export_knowledge_graph()
